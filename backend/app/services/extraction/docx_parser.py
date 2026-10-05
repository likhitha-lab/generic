"""DOCX extraction entrypoint.

Combines paragraphs, tables (table_parser), text boxes (textbox_parser),
and headers/footers into the final text, in a fixed order that matches how
a resume actually reads: name/contact in the header first, then the main
content, then any floating sidebar text boxes, then whatever's in the
footer last. This is the only module that knows that combined ordering -
table_parser/textbox_parser each only know their own piece.
"""
import io
import logging

from docx import Document
from docx.oxml.ns import qn

from app.core.config import settings
from app.services.extraction.table_parser import extract_tables_recursive
from app.services.extraction.textbox_parser import extract_textboxes
from app.services.extraction.xml_parser import paragraph_text, walk_block_container

logger = logging.getLogger(__name__)


def extract_paragraphs(doc: Document) -> list[str]:
    """Top-level body paragraph text, in order (a content-control <w:sdt>
    directly in the body, wrapping a paragraph, is unwrapped too). Tables
    are deliberately not included here - extract_tables_recursive owns
    those, so merge_document_text can control the combined ordering."""
    chunks: list[str] = []

    def walk(node) -> None:
        for child in node:
            if child.tag == qn("w:p"):
                text = paragraph_text(child).strip()
                if text:
                    chunks.append(text)
            elif child.tag == qn("w:sdt"):
                sdt_content = child.find(qn("w:sdtContent"))
                if sdt_content is not None:
                    walk(sdt_content)

    walk(doc.element.body)
    return chunks


def extract_headers(doc: Document) -> tuple[list[str], int]:
    """Every section's header content, deduped by element identity -
    sections marked "linked to previous" share the same underlying header
    part, so walking doc.sections naively would repeat identical header
    text once per section. Returns (chunks, header_count)."""
    chunks: list[str] = []
    seen: set[int] = set()
    header_count = 0
    for section in doc.sections:
        header_element = section.header._element
        if id(header_element) in seen:
            continue
        seen.add(id(header_element))
        header_chunks, _ = walk_block_container(header_element)
        if header_chunks:
            header_count += 1
        chunks.extend(header_chunks)
    return chunks, header_count


def extract_footers(doc: Document) -> tuple[list[str], int]:
    """Same as extract_headers, for footers. Returns (chunks, footer_count)."""
    chunks: list[str] = []
    seen: set[int] = set()
    footer_count = 0
    for section in doc.sections:
        footer_element = section.footer._element
        if id(footer_element) in seen:
            continue
        seen.add(id(footer_element))
        footer_chunks, _ = walk_block_container(footer_element)
        if footer_chunks:
            footer_count += 1
        chunks.extend(footer_chunks)
    return chunks, footer_count


def merge_document_text(doc: Document) -> tuple[str, dict]:
    """Combines every extractor above into the final text. Returns
    (full_text, stats) - stats feeds the DEBUG_RESUME_PIPELINE summary log
    in extract_docx_text below."""
    header_chunks, header_count = extract_headers(doc)
    paragraph_chunks = extract_paragraphs(doc)
    table_chunks, nested_table_count = extract_tables_recursive(doc)
    textbox_chunks, textbox_count = extract_textboxes(doc)
    footer_chunks, footer_count = extract_footers(doc)

    all_chunks = header_chunks + paragraph_chunks + table_chunks + textbox_chunks + footer_chunks
    full_text = "\n".join(all_chunks)

    stats = {
        "headers_found": header_count,
        "body_paragraphs": len(paragraph_chunks),
        "top_level_tables": len(doc.element.body.findall(qn("w:tbl"))),
        "nested_tables_found": nested_table_count,
        "textboxes_found": textbox_count,
        "footers_found": footer_count,
        "total_characters": len(full_text),
    }
    return full_text, stats


def extract_docx_text(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        # A malformed/corrupted DOCX (not a valid zip package, missing
        # required parts, etc.) would otherwise raise all the way up
        # through resume_service.py as an unhandled 500. Returning "" lets
        # the existing "no text extracted" check there turn this into the
        # same clean 422 a genuinely-empty file already gets, instead of
        # this module needing its own new exception type/handling upstream.
        logger.error("Failed to open this file as a DOCX (%s) - it may be corrupted or not a "
                     "real Word document despite the .docx extension.", exc)
        return ""

    full_text, stats = merge_document_text(doc)

    if settings.DEBUG_RESUME_PIPELINE:
        logger.info(
            "DOCX extraction complete: %(headers_found)d headers, %(body_paragraphs)d body "
            "paragraphs, %(top_level_tables)d top-level tables, %(nested_tables_found)d nested "
            "tables, %(textboxes_found)d text boxes, %(footers_found)d footers, "
            "%(total_characters)d total characters",
            stats,
        )
    if stats["nested_tables_found"]:
        logger.info(
            "DOCX had %d nested table(s) that Document.tables/_Cell.text alone would have "
            "missed entirely - now included.", stats["nested_tables_found"],
        )
    if stats["textboxes_found"]:
        logger.info(
            "DOCX had %d text box(es) that doc.paragraphs/doc.tables alone would have missed "
            "entirely - now included.", stats["textboxes_found"],
        )
    return full_text
