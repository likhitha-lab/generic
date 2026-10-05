"""Unit tests for file_generator.py's recruiter-controlled visibility
feature (show/hide email, phone, LinkedIn, address, employment dates) -
display-only: never mutates `parsed`, never removes data, respected by
both the PDF and DOCX renderers identically."""
import io

import pdfplumber
from docx import Document

from app.services.file_generator import (
    _apply_visibility_to_parsed,
    _contact_line,
    build_docx_bytes,
    build_pdf_bytes,
)

_PARSED = {
    "summary": "Experienced backend engineer with 8 years across cloud platforms.",
    "skills": ["Python", "AWS"],
    "tools": ["Git"],
    "education": ["B.Tech Computer Science"],
    "certifications": [],
    "experience": [
        {"company": "Acme Corp", "role": "Senior Engineer", "duration": "2016-2024",
         "points": ["Built the platform."], "is_career_break": False},
    ],
    "projects": [],
}
_NAME = "Jane Doe"
_CONTACT = {
    "email": "jane@example.com", "phone": "+1-555-0100",
    "location": "Austin, TX", "links": "linkedin.com/in/janedoe",
}


def _pdf_text(pdf_bytes: bytes) -> str:
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _docx_text(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    # Name/contact now live inside the header's name/logo table cells (see
    # _build_docx_with_tier) - document.paragraphs is body-only and never
    # includes header/footer content, and header.paragraphs alone misses
    # anything inside a header table, so all three have to be read here.
    header = document.sections[0].header
    header_text = "\n".join(p.text for p in header.paragraphs if p.text)
    header_table_text = "\n".join(
        p.text for table in header.tables for row in table.rows for cell in row.cells for p in cell.paragraphs if p.text
    )
    body_text = "\n".join(p.text for p in document.paragraphs if p.text)
    return header_text + "\n" + header_table_text + "\n" + body_text


# --- _contact_line -------------------------------------------------------------

def test_contact_line_shows_everything_by_default():
    line = _contact_line(_CONTACT)
    assert "jane@example.com" in line
    assert "+1-555-0100" in line
    assert "Austin, TX" in line
    assert "linkedin.com/in/janedoe" in line


def test_contact_line_hides_email_when_toggled_off():
    line = _contact_line(_CONTACT, {"show_email": False})
    assert "jane@example.com" not in line
    assert "+1-555-0100" in line  # everything else still shows


def test_contact_line_hides_phone_when_toggled_off():
    line = _contact_line(_CONTACT, {"show_phone": False})
    assert "+1-555-0100" not in line
    assert "jane@example.com" in line


def test_contact_line_hides_linkedin_when_toggled_off():
    line = _contact_line(_CONTACT, {"show_linkedin": False})
    assert "linkedin.com/in/janedoe" not in line
    assert "jane@example.com" in line


def test_contact_line_hides_address_when_toggled_off():
    line = _contact_line(_CONTACT, {"show_address": False})
    assert "Austin, TX" not in line
    assert "jane@example.com" in line


def test_contact_line_hides_all_when_all_toggled_off():
    line = _contact_line(_CONTACT, {
        "show_email": False, "show_phone": False, "show_linkedin": False, "show_address": False,
    })
    assert line == ""


def test_contact_line_none_visibility_means_show_everything():
    assert _contact_line(_CONTACT, None) == _contact_line(_CONTACT)


# --- _apply_visibility_to_parsed - never mutates the original ---------------

def test_apply_visibility_returns_same_object_when_dates_shown():
    result = _apply_visibility_to_parsed(_PARSED, None)
    assert result is _PARSED


def test_apply_visibility_blanks_duration_without_mutating_original():
    original_duration = _PARSED["experience"][0]["duration"]
    result = _apply_visibility_to_parsed(_PARSED, {"show_employment_dates": False})
    assert result is not _PARSED
    assert result["experience"][0]["duration"] == ""
    # The ORIGINAL dict is completely untouched - this is "hide only during
    # rendering, never remove from storage" at the data-structure level.
    assert _PARSED["experience"][0]["duration"] == original_duration


def test_apply_visibility_never_touches_other_experience_fields():
    result = _apply_visibility_to_parsed(_PARSED, {"show_employment_dates": False})
    entry = result["experience"][0]
    assert entry["company"] == "Acme Corp"
    assert entry["role"] == "Senior Engineer"
    assert entry["points"] == ["Built the platform."]


# --- End-to-end: PDF and DOCX both respect the same settings ---------------

def test_pdf_hides_employment_dates_when_toggled_off():
    pdf_bytes = build_pdf_bytes(_PARSED, _NAME, _CONTACT, visibility={"show_employment_dates": False})
    text = _pdf_text(pdf_bytes)
    assert "2016-2024" not in text
    assert "Acme Corp" in text  # company itself still shows


def test_docx_hides_employment_dates_when_toggled_off():
    docx_bytes = build_docx_bytes(_PARSED, _NAME, _CONTACT, visibility={"show_employment_dates": False})
    text = _docx_text(docx_bytes)
    assert "2016-2024" not in text
    assert "Acme Corp" in text


def test_pdf_and_docx_show_dates_identically_by_default():
    pdf_bytes = build_pdf_bytes(_PARSED, _NAME, _CONTACT)
    docx_bytes = build_docx_bytes(_PARSED, _NAME, _CONTACT)
    assert "2016-2024" in _pdf_text(pdf_bytes)
    assert "2016-2024" in _docx_text(docx_bytes)


def test_pdf_hides_contact_fields_matching_docx():
    visibility = {"show_email": False, "show_address": True, "show_phone": True, "show_linkedin": False}
    pdf_bytes = build_pdf_bytes(_PARSED, _NAME, _CONTACT, visibility=visibility)
    docx_bytes = build_docx_bytes(_PARSED, _NAME, _CONTACT, visibility=visibility)
    for text in (_pdf_text(pdf_bytes), _docx_text(docx_bytes)):
        assert "jane@example.com" not in text
        assert "linkedin.com/in/janedoe" not in text
        assert "+1-555-0100" in text
        assert "Austin, TX" in text


def test_hiding_dates_never_shrinks_page_cap_for_senior_candidate():
    # A candidate with 15+ years of real experience should still get the
    # wider page allowance (_max_pages_allowed) even when dates are hidden
    # from DISPLAY - page-fit estimation must read the TRUE duration, not
    # the display-filtered one.
    senior_parsed = dict(_PARSED)
    senior_parsed["experience"] = [
        {"company": f"Company {i}", "role": "Engineer", "duration": f"{2005 + i}-{2006 + i}",
         "points": ["Did work."], "is_career_break": False}
        for i in range(15)
    ]
    from app.services.file_generator import _max_pages_allowed
    assert _max_pages_allowed(senior_parsed["experience"]) == 5  # 15 years -> the 15+ bracket
    # Rendering with dates hidden must not raise and must still produce a
    # valid multi-entry PDF - a smoke test that the visibility-filtered
    # copy never reaches the page-cap calculation.
    pdf_bytes = build_pdf_bytes(senior_parsed, _NAME, _CONTACT, visibility={"show_employment_dates": False})
    text = _pdf_text(pdf_bytes)
    assert "Company 0" in text
    assert "2005-2006" not in text


def test_no_visibility_param_reproduces_exact_previous_behavior():
    # Backward compatibility - every caller before this feature existed
    # never passed `visibility` at all. Compares extracted TEXT, not raw
    # bytes - ReportLab embeds a fresh creation timestamp/ID on every
    # build, so two independent calls never produce byte-identical PDFs
    # even with identical content.
    with_none = build_pdf_bytes(_PARSED, _NAME, _CONTACT, visibility=None)
    without_param = build_pdf_bytes(_PARSED, _NAME, _CONTACT)
    assert _pdf_text(with_none) == _pdf_text(without_param)
